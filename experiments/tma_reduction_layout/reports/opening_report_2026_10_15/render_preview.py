"""Approximate geometry/text preview of the supplied PPTX; not PowerPoint rendering."""
from pathlib import Path
import math,json
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE,MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR,PP_ALIGN
from pptx.enum.dml import MSO_COLOR_TYPE,MSO_FILL_TYPE
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor

out=Path(__file__).resolve().parent
for name,path in [('CJK','/tmp/tma-ppt-fonts/NotoSC.ttf'),('CJKBold','/tmp/tma-ppt-fonts/NotoSCBold.ttf'),('Mono','/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf')]:
    pdfmetrics.registerFont(TTFont(name,path))
pres=Presentation(out/'tma_reduction_layout_opening_report_revised.pptx')
w,h=pres.slide_width/12700,pres.slide_height/12700
c=canvas.Canvas(str(out/'tma_reduction_layout_opening_report_preview.pdf'),pagesize=(w,h))
width_issues=[]
def rgb(color,default='171717'):
    try:
        if color.type==MSO_COLOR_TYPE.RGB: return str(color.rgb)
    except Exception: pass
    return default

def style(run,para):
    size=(run.font.size or para.font.size)
    size=size.pt if size else 18
    name=run.font.name or para.font.name or ''
    bold=run.font.bold if run.font.bold is not None else para.font.bold
    font='Mono' if name in ('Consolas','Courier New') else ('CJKBold' if bold else 'CJK')
    return size,font,rgb(run.font.color)

for page,slide in enumerate(pres.slides,1):
    c.setFillColor(HexColor('#ffffff'));c.rect(0,0,w,h,stroke=0,fill=1)
    for sh in slide.shapes:
        x,y,sw,shh=(float(getattr(sh,k))/12700 for k in ('left','top','width','height'))
        yy=h-y-shh
        if sh.shape_type==MSO_SHAPE_TYPE.LINE:
            c.setStrokeColor(HexColor('#'+rgb(sh.line.color,'C6C6C6')))
            c.setLineWidth(sh.line.width.pt if sh.line.width else 1)
            c.line(x,h-y,x+sw,h-y-shh)
            continue
        if sh.shape_type==MSO_SHAPE_TYPE.AUTO_SHAPE:
            fill=sh.fill.type==MSO_FILL_TYPE.SOLID
            line=sh.line.fill.type==MSO_FILL_TYPE.SOLID
            if fill: c.setFillColor(HexColor('#'+rgb(sh.fill.fore_color,'ffffff')))
            if line: c.setStrokeColor(HexColor('#'+rgb(sh.line.color,'C6C6C6')))
            c.setLineWidth(sh.line.width.pt if sh.line.width else 1)
            if sh.auto_shape_type==MSO_AUTO_SHAPE_TYPE.OVAL:
                c.ellipse(x,yy,x+sw,yy+shh,fill=int(fill),stroke=int(line))
            else: c.rect(x,yy,sw,shh,fill=int(fill),stroke=int(line))
        if not sh.has_text_frame: continue
        tf=sh.text_frame
        paragraphs=[]
        total=0
        for para in tf.paragraphs:
            runs=list(para.runs)
            if not runs: continue
            size,font,color=style(runs[0],para)
            spacing=para.line_spacing
            leading=spacing.pt if hasattr(spacing,'pt') else size*(spacing or 1.25)
            for text in para.text.replace('\v','\n').split('\n'):
                tw=pdfmetrics.stringWidth(text,font,size)
                available=sw-(tf.margin_left+tf.margin_right)/12700
                if tw>available+2:
                    width_issues.append({'page':page,'text':text,'width_points':round(tw,1),'available_points':round(available,1),'font_size':size,'wrap':tf.word_wrap})
                paragraphs.append((para,text,size,font,color,leading,tw))
                total+=leading
        if not paragraphs: continue
        top=y+tf.margin_top/12700
        usable=shh-(tf.margin_top+tf.margin_bottom)/12700
        if tf.vertical_anchor==MSO_ANCHOR.MIDDLE: top+=max(0,(usable-total)/2)
        elif tf.vertical_anchor==MSO_ANCHOR.BOTTOM: top+=max(0,usable-total)
        for para,text,size,font,color,leading,tw in paragraphs:
            xx=x+tf.margin_left/12700
            if para.alignment==PP_ALIGN.CENTER: xx=x+(sw-tw)/2
            elif para.alignment==PP_ALIGN.RIGHT: xx=x+sw-tf.margin_right/12700-tw
            c.setFont(font,size);c.setFillColor(HexColor('#'+color))
            c.drawString(xx,h-top-size,text)
            top+=leading
    c.showPage()
c.save()
(out/'preview_width_checks.json').write_text(json.dumps(width_issues,ensure_ascii=False,indent=2)+'\n')
print('Approximate preview saved; it does not substitute for native PowerPoint layout validation.')
print(json.dumps(width_issues,ensure_ascii=False,indent=2))
