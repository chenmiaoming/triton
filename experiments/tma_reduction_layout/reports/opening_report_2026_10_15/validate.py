"""Offline validation of the report and source/evidence bindings."""
import argparse
import hashlib
import io
import json
from pathlib import Path
import re
import subprocess
import tarfile
from zipfile import ZipFile

from pptx import Presentation
from audit import compute, read, sha

ROOT = Path(__file__).resolve().parent
REPO = ROOT.parents[3]


def validate(committed=False):
    result = compute()
    assert read(ROOT / 'audit.json') == result
    deck = ROOT / 'tma_reduction_layout_opening_report_revised.pptx'
    with ZipFile(deck) as z:
        assert z.testzip() is None
        assert not any(name.startswith('ppt/media/') for name in z.namelist())
    ppt = Presentation(deck)
    assert len(ppt.slides) == 20
    bodies, notes = [], []
    for page, slide in enumerate(ppt.slides, 1):
        bodies.append('\n'.join(sh.text for sh in slide.shapes if sh.has_text_frame))
        assert slide.has_notes_slide
        notes.append(slide.notes_slide.notes_text_frame.text)
        assert len(notes[-1]) > 80
        page_numbers = [sh.text for sh in slide.shapes if sh.has_text_frame
                        and sh.left / 914400 > 11.7 and sh.top / 914400 > 7.1 and sh.text.isdecimal()]
        assert page_numbers == [str(page)], (page, page_numbers)
        for sh in slide.shapes:
            assert sh.left >= 0 and sh.top >= 0
            assert sh.left + sh.width <= ppt.slide_width + 9144
            assert sh.top + sh.height <= ppt.slide_height + 9144
    all_body = '\n'.join(bodies)
    for wrong in ('36.0%', '41.2%', '36%–46%', '36%–41%', '模型低估', '模型高估 4 倍', '先转 FP32 再归约'):
        assert wrong not in all_body, wrong
    assert 'TMA 3.5.1' in bodies[5] and '早期 vec2 原型' in bodies[5]
    assert 'TMA 3.5.1' in bodies[6] and '早期 vec2 原型' in bodies[6]
    assert 'Global memory / cache' in bodies[4] and 'Shared memory' in bodies[4]
    assert 'min(E, 128 / element_bits)' in bodies[7]
    assert 'alignment' in bodies[7] and 'axis=1' in bodies[9]
    for number in ('298.368', '313.264', '297.984', '296.960', '298.880',
                   '422.400', '279.440', '285.696', '368.640', '421.888', '33.8%'):
        assert number in bodies[11], number
    summary = (ROOT.parents[1] / 'results/phase11/stage_b/summary.md').read_text()
    for condition in ('M32_N128_w4:bfloat16:max:host', 'M32_N64_w8:bfloat16:max:host'):
        line = next(line for line in summary.splitlines() if line.startswith('| ' + condition + ' |'))
        for number in [value.strip() for value in line.split('|')[2:6]]:
            assert number in bodies[12], (condition, number)
    assert '46.652%' in bodies[12] and '2.332%' in bodies[12]
    assert 'BF16 packed max 后转换' in bodies[13]
    assert all(term in bodies[14] for term in ('ML', 'MR', 'MO', 'MC', 'runtime K'))
    assert 'selection regret' in bodies[15] and '尚未执行' in notes[16]
    assert 'tl.load(x + offsets)' in bodies[17]
    assert '436.208' in bodies[18] and '不同批次' in bodies[18]
    assert len(read(ROOT / 'outline.json')) == 20
    assert read(ROOT / 'preview_width_checks.json') == []
    for check in read(ROOT / 'text_fit_checks.json'):
        assert check['used_height_in'] <= check['box_height_in'] + .025
    pdf = ROOT / 'tma_reduction_layout_opening_report_preview.pdf'
    info = subprocess.check_output(['pdfinfo', str(pdf)], text=True)
    assert int(re.search(r'^Pages:\s+(\d+)', info, re.M)[1]) == 20
    if committed:
        manifest = read(ROOT / 'deliverables_manifest.json')
        relative = ROOT.relative_to(REPO)
        archive = subprocess.check_output(['git', 'archive', 'HEAD', str(relative)], cwd=REPO)
        with tarfile.open(fileobj=io.BytesIO(archive)) as tar:
            assert tar.extractfile((relative / 'deliverables_manifest.json').as_posix()).read() == (
                ROOT / 'deliverables_manifest.json').read_bytes()
            for name, expected in manifest['files'].items():
                assert sha(ROOT / name) == expected, name
                content = tar.extractfile((relative / name).as_posix()).read()
                assert hashlib.sha256(content).hexdigest() == expected, name
    return {'status': 'PASS', 'slides': 20, 'main_slides': 17, 'appendix_slides': 3,
            'input_SHA256': result['input_attachment_SHA256'], 'PPTX_SHA256': sha(deck),
            'PDF_preview_SHA256': sha(pdf), 'source_and_evidence_bindings': 'PASS',
            'six_A10_and_two_H100_TTGIR_layouts': 'PASS', 'body_and_notes_checks': 'PASS',
            'text_fit_and_bounds': 'PASS', 'PDF_preview_pages': 20,
            'native_PowerPoint_rendered': False, 'new_GPU_experiments': 0}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--committed', action='store_true')
    args = parser.parse_args()
    result = validate(args.committed)
    (ROOT / 'validation.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    print('PASS: 20 slides, notes, eight existing layouts, source/evidence hashes, text bounds and PDF preview')
