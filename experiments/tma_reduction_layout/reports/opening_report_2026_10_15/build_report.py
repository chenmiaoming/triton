"""Revise the user's editable PPTX; preserve its layout diagrams and plain style.

No GPU, native build or experiment mutation. The PDF preview is reconstructed
from PPTX geometry using Noto SC; it is not a native PowerPoint export.
"""
import json
from pathlib import Path
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT = Path(__file__).resolve().parent
EXP = ROOT.parents[1]
AUDIT = json.loads((ROOT / 'audit.json').read_text())
INPUT = ROOT / 'input/original_opening_report.pptx'
OUTPUT = ROOT / 'tma_reduction_layout_opening_report_revised.pptx'
MAIN = AUDIT['source_versions']['observed_main']
PIN = AUDIT['source_versions']['experiment_upstream']
GH = f'https://github.com/triton-lang/triton/blob/{MAIN}/'
ISSUE = 'https://github.com/triton-lang/triton/issues/9703'

for name, file in (('CJK', '/tmp/tma-ppt-fonts/NotoSC.ttf'),
                   ('CJKBold', '/tmp/tma-ppt-fonts/NotoSCBold.ttf'),
                   ('Mono', '/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf')):
    pdfmetrics.registerFont(TTFont(name, file))

W, H = 13.33333224, 7.5
INK, MUTED, BLUE, LINE, PALE = '171717', '555555', '1D4569', 'C6C6C6', 'F4F4F4'


def lines_for(value, width, size, bold=False, mono=False):
    font = 'Mono' if mono else ('CJKBold' if bold else 'CJK')
    result = []
    for para in str(value).split('\n'):
        if not para:
            result.append('')
            continue
        if mono:
            assert pdfmetrics.stringWidth(para, font, size) <= width * 72, para
            result.append(para)
            continue
        tokens = re.findall(r'[A-Za-z0-9_./:+*<>=\-]+|\s+|.', para)
        merged = []
        for token in tokens:
            if merged and token in '。，；：！？、）】〉》…％%.,;:!?)]':
                merged[-1] += token
            else:
                merged.append(token)
        current = ''
        for token in merged:
            if pdfmetrics.stringWidth(current + token, font, size) <= width * 72 * .97:
                current += token
            else:
                if current:
                    result.append(current.rstrip())
                current = token.lstrip()
        result.append(current.rstrip())
    return result


class Author:
    def __init__(self):
        self.ppt = Presentation(INPUT)
        self.original = list(self.ppt.slides)
        self.order = []
        self.outline = []
        self.checks = []

    def text(self, slide, x, y, w, h, value, size=18, bold=False, color=INK,
             align='left', valign='top', mono=False, url=None):
        lines = lines_for(value, w, size, bold, mono)
        used = (size * 1.1 + (len(lines) - 1) * size * 1.23) / 72
        if used > h + .025:
            raise ValueError(f'Text overflow: {value!r}: {used:.3f} > {h}')
        assert min(x, y, w, h) >= 0 and x + w <= W + .01 and y + h <= H + .01
        sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = sh.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = False
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE if valign == 'middle' else MSO_ANCHOR.TOP
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_before = p.space_after = Pt(0)
            p.line_spacing = Pt(size * 1.23)
            p.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}[align]
            r = p.add_run()
            r.text = line
            r.font.name = 'Consolas' if mono else 'Microsoft YaHei'
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(color)
            ea = OxmlElement('a:ea')
            ea.set('typeface', 'Microsoft YaHei')
            r._r.get_or_add_rPr().append(ea)
            if url:
                r.hyperlink.address = url
        self.checks.append({'text': str(value), 'size': size, 'lines': len(lines),
                            'used_height_in': used, 'box_height_in': h})
        return sh

    def line(self, slide, x1, y1, x2, y2, color=LINE, width=1):
        sh = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        sh.line.color.rgb = RGBColor.from_string(color)
        sh.line.width = Pt(width)
        return sh

    def rect(self, slide, x, y, w, h, fill=PALE, stroke=None):
        sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(fill)
        if stroke:
            sh.line.color.rgb = RGBColor.from_string(stroke)
        else:
            sh.line.fill.background()
        return sh

    def box(self, slide, x, y, w, h, value, size=18, bold=False):
        self.rect(slide, x, y, w, h, 'FFFFFF', LINE)
        self.text(slide, x + .12, y + .04, w - .24, h - .08, value, size, bold,
                  align='center', valign='middle')

    def bullets(self, slide, x, y, w, items, size=18, gap=.20):
        for item in items:
            lines = lines_for(item, w - .28, size)
            height = (size * 1.1 + (len(lines) - 1) * size * 1.23) / 72 + .035
            dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + .035),
                                        Inches(y + size / 72 * .43), Inches(.055), Inches(.055))
            dot.fill.solid()
            dot.fill.fore_color.rgb = RGBColor.from_string(INK)
            dot.line.fill.background()
            self.text(slide, x + .28, y, w - .28, height, item, size)
            y += height + gap
        return y

    def table(self, slide, x, y, widths, rows, heights=.56, size=16, left_cols=(0,)):
        if isinstance(heights, (int, float)):
            heights = [heights] * len(rows)
        self.line(slide, x, y, x + sum(widths), y, INK, 1.1)
        yy = y
        for row_index, (row, height) in enumerate(zip(rows, heights)):
            xx = x
            for col, (width, value) in enumerate(zip(widths, row)):
                self.text(slide, xx + .10, yy + .055, width - .20, height - .10,
                          value, size, row_index == 0, align='left' if col in left_cols else 'center',
                          valign='middle')
                xx += width
            yy += height
            self.line(slide, x, yy, x + sum(widths), yy, LINE, .6)
        return yy

    def header(self, slide, section, title, subtitle='', source='', url=None):
        for sh in list(slide.shapes):
            sh._element.getparent().remove(sh._element)
        self.rect(slide, 0, 0, W, H, 'FFFFFF')
        self.text(slide, .60, .22, 11.7, .24, section, 10.5, color=MUTED)
        self.text(slide, .60, .63, 12.1, .58, title, 28, True)
        self.line(slide, .60, 1.27, 12.70, 1.27, BLUE, 1.2)
        if subtitle:
            self.text(slide, .60, 1.43, 12.10, .53, subtitle, 16, color=MUTED)
        self.line(slide, .60, 7.07, 12.70, 7.07, LINE, .5)
        self.text(slide, .60, 7.17, 11.30, .20, source, 9, color=MUTED, url=url)
        self.text(slide, 12.02, 7.17, .68, .20, '0', 10, color=MUTED, align='right')

    def page(self, section, title, subtitle='', source='', url=None, original=None):
        slide = self.original[original] if original is not None else self.ppt.slides.add_slide(self.ppt.slide_layouts[6])
        self.header(slide, section, title, subtitle, source, url)
        self.order.append(slide)
        self.outline.append({'title': title, 'section': section})
        return slide

    def keep(self, original, title):
        slide = self.original[original]
        self.order.append(slide)
        self.outline.append({'title': title, 'section': '保留并修订原图'})
        return slide

    def replace(self, slide, old, new, size=None):
        matches = [sh for sh in slide.shapes if sh.has_text_frame and sh.text == old]
        assert len(matches) == 1, (old, len(matches))
        sh = matches[0]
        p = sh.text_frame.paragraphs[0]
        run = p.runs[0]
        font_name, font_size, bold = run.font.name, run.font.size, run.font.bold
        color = str(run.font.color.rgb)
        alignment = p.alignment
        sh.text_frame.clear()
        p = sh.text_frame.paragraphs[0]
        p.alignment = alignment
        p.space_before = p.space_after = Pt(0)
        r = p.add_run()
        r.text = new
        r.font.name = font_name
        r.font.size = Pt(size) if size else font_size
        r.font.bold = bold
        r.font.color.rgb = RGBColor.from_string(color)
        return sh

    def note(self, slide, value):
        slide.notes_slide.notes_text_frame.text = value

    def save(self):
        ids = self.ppt.slides._sldIdLst
        by_id = {node.id: node for node in ids}
        assert len(self.order) == len(ids) == 20
        ordered_ids = [by_id[slide.slide_id] for slide in self.order]
        for node in list(ids):
            ids.remove(node)
        for number, (slide, node) in enumerate(zip(self.order, ordered_ids), 1):
            ids.append(node)
            for sh in slide.shapes:
                if sh.has_text_frame and sh.left > Inches(11.7) and sh.top > Inches(7.10) and sh.text.isdecimal():
                    sh.text_frame.paragraphs[0].runs[0].text = str(number)
            self.outline[number - 1]['page'] = number
        self.ppt.core_properties.title = '面向访存与部分归约的 GPU 张量布局选择'
        self.ppt.core_properties.subject = '问题、当前编译器选择机制、已有实验与代价建模计划'
        self.ppt.core_properties.comments = 'Revised from the user-provided deck; evidence and source audit frozen 2026-10-15.'
        self.ppt.save(OUTPUT)
        (ROOT / 'outline.json').write_text(json.dumps(self.outline, ensure_ascii=False, indent=2) + '\n')
        (ROOT / 'text_fit_checks.json').write_text(json.dumps(self.checks, ensure_ascii=False, indent=2) + '\n')
        notes = []
        for number, slide in enumerate(self.order, 1):
            notes.append(f'## 第 {number} 页：{self.outline[number-1]["title"]}\n\n' + slide.notes_slide.notes_text_frame.text)
        (ROOT / 'speaker_notes.md').write_text('\n\n'.join(notes) + '\n')


def build():
    a = Author()

    s = a.keep(0, '面向访存与部分归约的 GPU 张量布局选择')
    a.replace(s, 'GPU 编译器张量布局（Layout）自动选择研究', '面向访存与部分归约的 GPU 张量布局选择')
    a.replace(s, '面向 Triton 访存—归约（Load–Reduction）冲突的机理分析与代价建模',
              '以 Triton 为切入点：问题表征、现有编译器机制与代价建模方案')
    a.replace(s, '开题方向定位、核心问题机理与代价模型路线', '科研进展、现有选择机制与后续研究计划')
    a.replace(s, '开题背景与定位 → 归约计算与硬件路径 → 布局冲突机理 → 初版模型反思与新方案',
              '计算语义 → 布局与现有选择机制 → 受控实验 → 代价建模计划')
    a.note(s, '本次汇报已完成的问题表征、源码与编译产物审计、Gluon 受控布局对照和本地 A10 版本对照。'
           '新微基准 C1、正式硬件校准、静态选择模型和新留出集验证均尚未执行。'
           '研究目标是识别值得换布局的部分规约配置，同时保留默认布局已经较好的情况。'
           '\n源码版本：固定 upstream ' + PIN + '；本地 Triton 3.7.1；本次观察 main ' + MAIN + '。')

    s = a.page('1. 研究问题', '切入点：取数布局与部分归约的局部性',
               '从 TMA 的版本回归出发，研究哪些配置需要改选布局，哪些配置应保留默认。',
               '[1] Issue #9703；[2,3] H100 与 A10 既有实验', ISSUE, original=1)
    a.bullets(s, .82, 2.12, 11.65, [
        '编译器需要把一个逻辑张量分配到寄存器、lane 和 warp；同一数学计算可以有不同布局。',
        '连续 N 维的向量化可能降低取数成本；沿 M 维规约的线程局部性可能减少通信。',
        '已有 BF16/max 对照表明：某些配置有明显优化空间，另一些配置的默认布局已接近较快候选。',
    ], 19, .24)
    a.line(s, .82, 4.62, 12.48, 4.62, LINE)
    a.text(s, .84, 4.90, 11.60, .40, '研究问题', 20, True)
    a.bullets(s, .86, 5.47, 11.55, [
        '如何用编译期信息和硬件校准参数预测完整“取数—归约—输出”的成本？',
        '能否在未见形状上选到较好布局，并在收益小或预测不确定时回退默认？',
    ], 19, .18)
    a.note(s, '不要把“默认布局存在更快候选”与“某版本升级后发生回归”混为一谈。'
           'H100 TMA 与 A10 普通 load 的软件版本和测量批次不同，分别报告，不用一个收益区间概括。'
           'vec 只是布局的一部分；现阶段尚未分离各项硬件成本，也没有经过校准的选择器。')

    s = a.original[2]
    a.order.append(s)
    a.outline.append({'title': '代表性 Triton 程序：TMA 后逐列 max', 'section': '1. 研究问题'})
    a.replace(s, '代表性算子实例：矩阵分块读取后执行逐列归约', '代表性 Triton 程序：TMA 后逐列 max')
    a.replace(s, '输入 BF16 [1,8192,8192]；每块读取 [32,128] 或 [32,64]；沿行方向 M（axis=1）求每列最大值。',
              '输入 BF16 [1,8192,8192]；加载 [1,BM,BN]，沿 M 维（axis=1）规约，输出 [1,BN]。')
    a.replace(s, '共性现象：无论 Hopper TMA 加载还是 A10 普通 tl.load，编译器默认均选 vec8 布局。',
              '本页展示 TMA 路径；x_desc 在 host 端构造，descriptor 构造不在 kernel 计时区间。', 15)
    a.replace(s, '实测差距：32×64 (8 warps) 下改为按列集中布局，H100 耗时降 41.2%，A10 降 36.0%。',
              '原始几何：32×128 / 4 warps；补充几何：32×64 / 8 warps。普通 tl.load 完整代码见附录 A。', 15)
    a.replace(s, '[1] 原始 Issue #9703 复现程序；[2] A10 普通 tl.load / cp.async 对照实验',
              '[1] 原始 Issue #9703；host-created descriptor，BF16 partial max')
    a.note(s, '代码保留三维 tile [1,BM,BN]，所以 axis=1 是 M，而不是二维代码中的 axis=0。'
           'host 端构造 TensorDescriptor(x, list(x.shape), list(x.stride()), [1,BM,BN])。'
           '原始配置启动 16384 个 CTA，num_warps=4；每个 tile 的部分结果写入 [256,8192] FP32 输出。'
           '该页只说明 TMA 程序，不能把它当作 A10 上执行了 TMA 的证据。普通 load 代码见第18页。')

    s = a.keep(3, '计算模式：逐列归约，消去 M、保留 N')
    a.replace(s, '用 4×4 矩阵说明语义；实际案例为 32×128 的输入 tile。',
              '用 4×4 矩阵说明一个 tile 内的计算；全局 M 维最大值还需要合并不同 tile 的部分结果。')
    a.note(s, '每列的 BM 个元素合并成一个最大值，消去 M，保留 N。矩阵例子的输出 5、9、8、7 正确。'
           '这是分块 partial reduction，不是一次完成整个 M=8192 的最终列最大值。'
           '后面改变的是元素的硬件归属，数学语义保持不变。')

    s = a.page('1. 研究问题', '两种取数路径，归约算术都使用寄存器中的值',
               'TMA 的初始 shared staging，与归约/输出转换使用的 shared 临时存储，需要分开。',
               '[1,4] Triton/Gluon 数据路径；[3] A10 归档 SASS 与 NCU', original=4)
    a.text(s, .80, 2.03, 11.7, .30, 'TMA / host descriptor', 18, True)
    for x, label in ((.85, 'Global memory'), (4.00, 'Shared memory'), (7.15, 'Registers')):
        a.box(s, x, 2.48, 2.55, .54, label, 19)
    a.text(s, 3.46, 2.53, .40, .35, '→', 20, align='center')
    a.text(s, 6.60, 2.53, .40, .35, '→', 20, align='center')
    a.text(s, 3.36, 2.10, .65, .27, 'TMA', 14, align='center')
    a.text(s, 6.48, 2.10, .95, .27, 'local_load', 14, align='center')
    a.text(s, 10.03, 2.55, 2.42, .35, '→  partial max', 19)
    a.text(s, .80, 3.45, 11.7, .30, '普通 tl.load', 18, True)
    a.box(s, .85, 3.95, 3.50, .54, 'Global memory / cache', 19)
    a.text(s, 4.60, 4.00, 2.10, .35, '→  tl.load  →', 18, align='center')
    a.box(s, 7.15, 3.95, 2.55, .54, 'Registers', 19)
    a.text(s, 10.03, 4.02, 2.42, .35, '→  partial max', 19)
    a.bullets(s, .86, 5.13, 11.55, [
        '跨线程规约或结果重排仍可能使用 shared memory 和同步；它们不是普通 load 的初始取数。',
        'vec = sizePerThread[N] 描述基本连续元素归属；不是 TMA 传输宽度，也不等于硬件加载次数。',
    ], 18, .24)
    a.note(s, '原普通 load 诊断的 direct vec1 在 32×128/w4 上没有 LDS/STS/SHFL/BAR，max 直接在寄存器上执行；'
           'direct vec8 的 LDG 也直接进入寄存器，但后续规约/转换有 shared 操作。cp.async 是另一条显式 shared staging 对照，'
           '不能把它叫 TMA。这里的 max 是编译器生成的寄存器算术；没有声称在 global memory 中直接做规约。')

    s = a.keep(5, '三种输入布局：参数、重复覆盖与贡献线程')
    a.replace(s, '3.5.1（报告）', 'TMA 3.5.1（报告）', 13.5)
    a.replace(s, '3.6.0（报告）', 'TMA 3.6.0（报告）', 13.5)
    a.replace(s, '当前原型', '早期 vec2 原型', 13.5)
    a.replace(s, '[1,3] 历史 encoding 与布局定义；[4] 当前原型 TTGIR',
              '[1,4,5] 原 issue 的 TMA encoding、布局定义与早期原型 TTGIR')
    a.replace(s, '重复覆盖使用同一组线程；多行通常不连续。当前 baseline 的输入布局也为 vec8。',
              '重复覆盖不新增线程。第三列仅作布局示例；早期选择规则已弃用，新模型尚待建立。', 14)
    a.note(s, '本表所有参数、基本覆盖、重复次数和贡献线程数均经复核。'
           '历史两列是原 issue 的 TMA encoding，不是普通 load 在旧版本的 encoding。'
           '普通 load 的被测 3.5.1/3.6.0/3.7.1 都已选择 vec8。第三列是 Phase10 的早期 vec2 原型，只用于说明不同 ownership，'
           '不代表已验证的现行成本模型。N 优先时 vec2 的列0属于 warp0 和 warp2。')

    s = a.keep(6, '第 0 列的 32 个值分别归哪个线程？')
    a.replace(s, '3.5.1（原报告）', 'TMA 3.5.1（原报告）', 12.5)
    a.replace(s, '3.6.0 / 当前 baseline 输入', 'TMA 3.6.0 / 固定 baseline', 12.5)
    a.replace(s, '当前原型：upstream + 补丁', '早期 vec2 原型：布局示例', 12.5)
    a.note(s, '保持用户原先要求的完整32行归属图。TMA旧布局每列1个线程；vec8每列8个线程跨4 warp；早期vec2每列2个线程跨2 warp。'
           'vec8的同一W0:L0持有M=0/8/16/24，完整tile上各持有N=0..7，共32元素。'
           'vec2的第0列偶数行是W0:L0、奇数行W2:L0。此图只是输入元素归属，不表示实际指令顺序或已证明的时间。'
           '行标签约9pt，现场可放大讲解；具体数学不变。')

    s = a.page('2. 当前编译器选择机制', 'TMA 与普通 load：两个入口，都是先构造访存布局',
               '核对实验 upstream 20d9d98、v3.7.1 与今日 main 4eda6cc；核心选择规则一致，非新增性能实验。',
               '[S1,S2] Coalesce.cpp / CoalesceUtils.cpp / Utility.cpp',
               GH + 'lib/Dialect/TritonGPU/Transforms/Coalesce.cpp')
    a.text(s, .82, 2.06, 5.60, .38, 'TMA / descriptor', 21, True)
    a.text(s, 6.85, 2.06, 5.60, .38, '普通 tl.load', 21, True)
    a.line(s, 6.61, 2.02, 6.61, 5.91, LINE, .7)
    a.text(s, .84, 2.63, 5.35, 1.0,
           '# pickDescriptorLoadStoreLayout\nE = max(numElems / (warps * 32), 1)\nV = min(E, 128 / element_bits)\nSPT = [1, ..., V]', 13, mono=True)
    a.text(s, 6.88, 2.63, 5.40, 1.0,
           '# buildCoalescedEncoding\norder <- pointer contiguity\nV0 <- min(alignment, contiguity,\n          128 / element_bits)', 13, mono=True)
    a.bullets(s, .84, 3.94, 5.35, [
        'N 优先；BF16 的目标上限为 8，FP32 为 4。',
        '再按 shape 和 warp 数分配 lane/warp。',
        '此函数不读取后续 reduce axis。',
    ], 17, .20)
    a.bullets(s, 6.88, 3.94, 5.39, [
        '用指针连续性与对齐约束向量宽度。',
        '协调相关同 shape/order 的内存操作，并受每线程元素量限制。',
        '再调用同一个 blocked-layout 构造器。',
    ], 17, .20)
    a.text(s, .84, 6.30, 11.60, .48,
           '这两个入口没有对本例的 vec1/2/4/8 做完整“加载—部分规约—输出”的校准时间比较。', 18, True)
    a.note(s, 'E中的除法是源码的整除。128单位是bit，而不是bytes或TMA交易大小。'
           '普通load的简式在单一BF16连续指针场景下解释，严格源码还在相关同shape同order内存操作间取最大值，再受E限制。'
           '不推广到atomic特殊限制，不把SPT当动态load次数。两条路径的源码入口不同，但最终都构造BlockedEncodingAttr。'
           'source_audit中固定保存今日main SHA，关键五个transform文件与实验upstream逐字相同；没有计时新main。')

    s = a.page('2. 当前编译器选择机制', '同样选 vec8，为什么同列贡献线程会从 8 变成 32？',
               'BlockedEncodingAttr 的构造器从连续 N 维开始分配 lane/warp，再把剩余资源放到 M 维。',
               '[S3] TritonGPUAttrDefs.td：shape-based builder；[3] 实际 TTGIR',
               f'https://github.com/triton-lang/triton/blob/{PIN}/include/triton/Dialect/TritonGPU/IR/TritonGPUAttrDefs.td#L827-L863')
    a.table(s, .78, 2.14, [3.30, 4.45, 4.15], [
        ['分配过程', '32×128 / 4 warps', '32×64 / 8 warps'],
        ['平均元素 / 线程', '4096 / 128 = 32', '2048 / 256 = 8'],
        ['默认 V（BF16）', 'min(32, 8) = 8', 'min(8, 8) = 8'],
        ['N 上的线程位置', '128 / 8 = 16', '64 / 8 = 8'],
        ['threadsPerWarp [M,N]', '[2,16]', '[4,8]'],
        ['warpsPerCTA [M,N]', '[4,1]', '[8,1]'],
        ['每列贡献线程', '2 × 4 = 8', '4 × 8 = 32'],
    ], .53, 17)
    a.text(s, .85, 6.20, 11.60, .55,
           '规则使用了 shape 和 warp 数，但没有用部分规约的预测时间来决定如何分配这些线程。', 18, True)
    a.note(s, 'N128/vec8只需16个线程位置，因此一个warp的剩余2个lane位置用于M；4个warp全部用于M。'
           'N64/vec8只需8个位置，剩余4个lane位置用于M；8个warp用于M，因此同一列有32个贡献线程。'
           '输入布局与后续lowering成本并非一一对应，不能仅靠贡献线程数预测总时间。'
           'audit.py在CPU复现builder，对照六个已存release/shape TTGIR全部一致；不是新增instrumented compiler trace。')

    s = a.page('2. 当前编译器选择机制', '后续 pass 为什么没有把本例改回线程局部布局？',
               '编译器有布局传播与局部优化；本例没有命中能改写它的特定规约模式。',
               '[S4,S5] RemoveLayoutConversions.cpp / OptimizeThreadLocality.cpp',
               GH + 'lib/Dialect/TritonGPU/Transforms/OptimizeThreadLocality.cpp#L254-L314')
    for x, text in ((.82, 'Coalesce'), (4.76, 'RemoveLayoutConversions'), (8.70, 'OptimizeThreadLocality')):
        a.box(s, x, 2.13, 3.55, .59, text, 16, True)
    a.text(s, 4.40, 2.26, .27, .31, '→', 17)
    a.text(s, 8.34, 2.26, .27, .31, '→', 17)
    a.table(s, .82, 3.10, [3.30, 8.30], [
        ['相关阶段', '本例中的行为 / 限制'],
        ['布局传播', 'descriptor 是 anchor；普通 load 在本例也作为内存 anchor，布局可向 cast / reduce 等传播。'],
        ['规约模式优化', '循环规约变换要求 axis == rank−1，并匹配特定 scf.for 累加结构等条件。'],
        ['本例', '[1,BM,BN] 的 axis=1，不是最后一维；也没有对应循环结构，因此不匹配。'],
    ], [.55, .87, .87, .80], 17, (0, 1))
    a.text(s, .86, 6.43, 11.55, .43,
           '最终 TTGIR 保留 vec8；这说明规则链的结果，不说明编译器比较过所有候选的完整成本。', 17, True)
    a.note(s, '图为相关pass的简化关系，实际pipeline多次运行RemoveLayoutConversions，且包含其他pass与TMA lowering。'
           'OptimizeThreadLocality还处理特定reshape/gather；本例没有归约前允许重排的reshape。'
           '所核对版本没有名为shouldOptimizeThreadLocality的函数，具体guards位于runOnOperation。'
           'RemoveLayoutConversions也有rematerialization局部代价估计，因此不说整个编译器没有成本机制。'
           '缺口是对本例候选输入布局的完整、经过校准的联合时间比较。')

    s = a.page('2. 当前编译器选择机制', '需要比较完整路径，同时满足资源约束',
               '同列贡献线程少、vec 更小、静态指令更少，都不是完整 kernel 更快的充分条件。',
               '[1,S1–S5] 问题讨论、源码与已有编译产物', original=7)
    a.table(s, .83, 2.12, [3.15, 4.50, 4.15], [
        ['成本 / 约束', '布局改变的内容', '建模要求'],
        ['取数路径', 'TMA shared-load；普通 global/cache load', '按实际访问模板校准'],
        ['线程内规约', '每线程同列值数、packed max / widening', '预测实际操作及依赖'],
        ['跨线程合并', 'lane / warp 归属、shuffle / shared / barrier', '校准通信与同步子图'],
        ['输出路径', '结果 SliceLayout、转换与 store', '保留实际消费者要求'],
        ['资源可行性', '寄存器、SMEM、spill 与可达驻留', '先合法性门，再预测时间'],
    ], [.55, .72, .72, .72, .65, .65], 16, (0, 1, 2))
    a.text(s, .86, 6.45, 11.55, .40,
           '目标是预测成本与适用条件；用微基准校准参数，再用完整 kernel 检验组合误差。', 17)
    a.note(s, '现阶段已经看到不同lowering、转换数与通信结构，但尚未把总耗时可靠分解为独立分项。'
           '资源变化影响吞吐和延迟隐藏，不能把所有阶段时间简单相加，也不能把两个不同kernel的耗时差当纯规约延迟。')

    s = a.page('3. 已完成实验', 'A10 普通 load：默认布局与固定候选的对照',
               'BF16/max；同批 Triton 3.7.1；单位 µs。普通 Triton 与显式布局 Gluon 分别保留产物。',
               '[3] results/local_a10/20261015_gpu0_r2：3 次进程重复，完整输出检查通过')
    rows = [['tile / warps', 'Triton 默认', 'Gluon vec1', 'vec2', 'vec4', 'vec8']]
    for case in AUDIT['cases']:
        t = case['A10_same_cohort_us']
        rows.append([f'32×{case["BN"]} / {case["warps"]}',
                     f'{t["canonical"]:.3f}', f'{t["direct:vec1"]:.3f}', f'{t["direct:vec2"]:.3f}',
                     f'{t["direct:vec4"]:.3f}', f'{t["direct:vec8"]:.3f}'])
    a.table(s, .82, 2.22, [2.10, 2.00, 2.15, 1.82, 1.82, 1.82], rows, .72, 17)
    a.bullets(s, .87, 4.87, 11.55, [
        '32×128 / w4：默认实现接近已测较快候选；直接改成 vec1 反而更慢。',
        '32×64 / w8：同一 Gluon kernel 的 vec8→vec1，耗时减少约 33.8%（约 1.51× 加速）。',
    ], 19, .26)
    a.text(s, .87, 6.42, 11.55, .43,
           '描述性对照，不宣称全局或唯一最优；Triton↔Gluon 的时间差不直接作为纯布局因果收益。', 15.5)
    a.note(s, '所有时间为同批三个worker中位数的中位数。32×64/w8原始Triton默认422.400；Gluon vec8 421.888、vec1 279.440。'
           '33.8%=1−279.440/421.888，而不是把另一轮版本实验的436µs与这里279µs混算。'
           '32×128/w4还测了vec8 within-warp等变体，详见完整证据。接近已测最快候选是描述性观察，不等于全空间最优。'
           '计时与NCU分开；既有NCU不能直接把每项总性能差归因于单一机制。')

    s = a.page('3. 已完成实验', 'H100 TMA：收益同样依赖形状与 warp 数',
               '固定未修改 upstream 20d9d98；BF16/max；原始 host descriptor 路径；full grid。',
               '[2] Phase 11 Stage B interpretation.md / summary.md：7 次 worker 调用')
    a.table(s, .82, 2.14, [2.65, 2.15, 2.15, 2.15, 2.15], [
        ['tile / warps', 'vec1 µs', 'vec2 µs', 'vec4 µs', 'vec8 µs'],
        ['32×128 / 4', '56.048', '56.160', '56.464', '57.296'],
        ['32×64 / 8', '61.392', '62.048', '68.528', '112.096'],
    ], .62, 17)
    a.table(s, .82, 4.40, [2.65, 3.55, 5.05], [
        ['配对 vec1 / vec8 对照', '耗时减少估计', '95% 同条件同时区间'],
        ['32×128 / 4', '2.332%', '[0.801%, 3.841%]'],
        ['32×64 / 8', '46.652%', '[41.056%, 51.718%]'],
    ], .53, 16)
    a.text(s, .87, 6.31, 11.55, .51,
           '原始几何差异小；补充几何有明确机会，vec1/vec2 尚未分出超过 3% 的稳定差异。', 17)
    a.note(s, '第一表是worker中位数的中位数；百分比来自七次worker配对log-ratio，使用同条件六个比较的Bonferroni/t区间。'
           '因此不能直接用第一表两个跨worker中位数相除复算百分比。3%是预声明实用差异阈值。'
           '这里固定host路径；device descriptor的41.297%保存在历史结果，不混入原始路径汇报。'
           '已有默认Triton↔matching-layout Gluon bridge证据支持入口对应；这仍不是新compiler patch已实现这些收益的证明。'
           '各候选资源与lowering共同变化尚需后续机制校准。')

    s = a.page('4. 初版模型反思', '早期模型的不足：逻辑分数不能替代硬件成本',
               '以下是早期 H100 32×128/w4 归档代码的静态对照，不是本轮测得的阶段时间。',
               '[5] Phase 10 归档 PTX；[6] NCU 指标定义', original=8)
    a.table(s, .83, 2.10, [4.00, 3.90, 3.90], [
        ['项目', '默认 vec8', '早期 vec2 候选'],
        ['旧加载分数（无量纲，已弃用）', '4', '16'],
        ['初始加载 PTX 静态条数', '4', '4'],
        ['PTX opcode', 'ld.shared.v4.b32', 'ldmatrix...x4...b16'],
        ['BF16→FP32 转换静态条数', '8', '2'],
        ['局部操作顺序', 'BF16 packed max 后转换', 'BF16 packed max 后转换'],
        ['初始加载硬件成本', '此对照未单独确认', '此对照未单独确认'],
    ], [.54, .58, .54, .58, .54, .66, .54], 15.5, (0,))
    a.text(s, .87, 6.36, 11.55, .45,
           '静态结构需要匹配实际 lowering，再经微基准校准；PTX 条数不能换算 requests、wavefront 或时间。', 16.5)
    a.note(s, '默认PTX 120行起先执行max.bf16x2，127行起执行cvt.f32.bf16；候选127行起packed max、143行起转换。'
           '因此两者都是先BF16局部合并再转换，而不是只有vec2能延后转换。'
           '旧加载分4/16与两种不同指令静态条数4/4没有可直接换算的物理单位，删除了“模型低估/高估4倍”。'
           '两份代码最初加载段的4条PTX不能证明硬件成本相等。旧等权规则已弃用，不安排继续证明其正确性或错误性。')

    s = a.page('5. 下一步方案', '先验证微基准能测到什么，再校准子图有效成本',
               '首版固定 H100、host TMA、BF16 partial max；循环、反馈和 checksum 的附加工作显式保留。',
               '[7] Phase 11 FOLLOWUP_PLAN.md：review draft，尚未执行 C1', original=9)
    a.table(s, .82, 2.10, [1.40, 6.00, 4.40], [
        ['子图', '合法 Gluon 构造', '用途'],
        ['ML 取数', 'TMA 预填 shared；指定布局读取并形成 checksum', '访问模板、有效延迟/吞吐'],
        ['MR 规约', '驻留寄存器；max 与有界、非幂等的数据依赖反馈', '局部计算与交换/同步子图'],
        ['MO 输出', '实际结果布局 → 固定输出布局 → store', '消费者与转换路径成本'],
        ['MC 全图', '原始 host TMA → load → max → 输出', '检验组合预测与资源匹配'],
    ], [.54, .68, .79, .68, .68], 16, (0, 1, 2))
    a.bullets(s, .88, 5.78, 11.5, [
        '同一 CUBIN 用 runtime K=1/16/64/256；NCU 核对目标操作确实随 K 增长。',
        'C1 最多 16 个程序，先检查重复、模板与资源；无法辨识时修订方案，不冒充裸指令延迟。',
    ], 16.5, .16)
    a.note(s, 'ML/MR需要可观察的结果和合法block-uniform runtime循环；不能重复max(x)同一个不变张量就当重复执行。'
           '人工反馈成本不直接转嫁到原kernel；不能简单相减得纯max或纯转换延迟。MR若与原kernel的BF16 packed模板不匹配，'
           '只能校准匹配部分。MO若转换被折叠，只保留诊断并使用真实相邻子图联合辨识。'
           'D=1/2和驻留限制区分依赖延迟与吞吐。纯寄存器子图不需要descriptor。'
           '将来扩展普通load时需单独的global/cache取数模板，不能直接套shared-load参数。')

    s = a.page('5. 下一步方案', '模型在编译期做选择，完整 kernel 的实测用于评价',
               '灰盒：静态依赖与资源结构 + 每架构校准参数；不是逐个运行待优化 kernel 再挑赢家。',
               '[7] FOLLOWUP_PLAN.md：静态计划、参数辨识、冻结预测与留出评价')
    a.box(s, .82, 2.14, 3.55, .86, '编译期特征 / target\nshape / dtype / layout', 17, True)
    a.box(s, 4.76, 2.14, 3.55, .86, '依赖图与资源计划\n硬件校准参数 θ', 17, True)
    a.box(s, 8.70, 2.14, 3.55, .86, '预测各合法候选时间\n不确定则回退默认', 17, True)
    a.text(s, 4.40, 2.45, .27, .31, '→', 17)
    a.text(s, 8.34, 2.45, .27, .31, '→', 17)
    a.bullets(s, .88, 3.57, 11.55, [
        '预测时只用静态信息与已冻结参数；测试 kernel 的实测时间、NCU 与最终寄存器数不作输入。',
        '对照默认布局、固定 vec 与简单局部性基线；报告收益/退化、回退覆盖率与最坏选择损失。',
        '模型冻结后测新的留出形状；跨架构区分参数直接迁移与仅重校准硬件参数的迁移。',
    ], 18, .20)
    a.rect(s, .86, 6.05, 11.60, .63)
    a.text(s, 1.04, 6.20, 11.21, .34,
           'selection regret = T(selected) / min(T(same candidate set)) - 1', 15.2, mono=True)
    a.note(s, '这是待验证模型路线，不是已实现选择器。θ应具有延迟或服务速率等可解释单位，不给每个layout分别自由拟合权重。'
           '寄存器与驻留使用早期静态估计和不确定区间，不能偷用测试binary的真实资源信息。'
           'regret中的实测最小值是有噪声的评价oracle，需报告不确定性。'
           '首版优先验证H100未见shape，之后才考虑B200与其他可用架构；旧数据用于调优后不能继续称held-out。')

    s = a.page('6. 研究进度与计划', '已完成问题表征；下一步是校准可行性与选择模型',
               '把已经完成的证据、尚待验证的机制，以及后续交付分开汇报。',
               '[2,3,7] 既有实验、源码审计与后续 review draft', original=10)
    a.table(s, .82, 2.12, [2.15, 6.15, 3.50], [
        ['阶段', '具体内容 / 可交付', '状态'],
        ['问题表征', 'TMA/Gluon 与 A10 对照、默认选择机制、原始证据', '已完成；分项归因待校准'],
        ['C1 可行性', '8 布局的 ML/MR，检查动态重复、模板与资源', '计划中；最多 16 个程序'],
        ['C2 / C3', '正式校准、静态计划、参数辨识与冻结选择器', 'C1 通过后推进'],
        ['D1 / D2', '新形状留出评价；H100 后再做 B200 等迁移', '模型冻结后执行'],
    ], [.55, .79, .79, .79, .79], 16, (0, 1, 2))
    a.text(s, .87, 6.22, 11.55, .60,
           '预期产出：可解释的布局选择模型、可复现实验与编译器原型；上游改动以验证结果为依据。', 18)
    a.note(s, '实际进度为Phase11 A/B、两轮本地A10与本次源码/文档审计已完成；新的C1尚未执行，模型尚未构建。'
           '不把早期Phase10的StageC与这次新的校准C1混淆。'
           '目前已有显著优化机会和负方向对照，但没有把所有机理项测清，也没有证明跨架构泛化。'
           '创建PR仍需用户另行许可；此次任务只更新报告和整理代码机制。')

    s = a.page('附录 A. 普通 load 实际程序', 'A10 上运行的普通 Triton DSL',
               '与前述 TMA 程序具有相同数学操作；仍保留 [1,BM,BN]，所以规约 axis=1。',
               '[3] local_a10_versions/kernels.py：与上一轮 canonical 函数 AST 一致')
    code = (EXP / 'local_a10_versions/kernels.py').read_text()
    code = code[code.index('@triton.jit'):].rstrip()
    a.rect(s, .83, 2.08, 11.65, 3.85)
    a.text(s, 1.02, 2.25, 11.25, 3.48, code, 13.2, mono=True)
    a.text(s, .87, 6.26, 11.55, .52,
           '输入 [1,8192,8192] BF16；输出 [256,8192] FP32；未在源码指定布局，由编译器自动选择。', 17)
    a.note(s, '该函数源码与已完成A10试验中的普通Triton anchor一致；新报告未新编译或执行GPU。'
           'BM=32，BN128/w4的grid=16384；BN64/w8的grid=32768。'
           '没有descriptor，没有手工shared allocation，也没有Gluon显式布局；shared通信由后端按实际规约/转换需求生成。')

    s = a.page('附录 B. 版本回归与优化空间', 'A10：没有版本回归，不代表默认布局总是合适',
               '独立版本对照批次；同 Python/Torch/ptxas；统一加载归档 CUBIN 配对计时。',
               '[3] local_a10_versions/20261015_gpu0_r2：7 次进程重复，2,520 个 event 样本')
    a.table(s, .82, 2.19, [3.10, 2.90, 2.90, 2.90], [
        ['配置', '3.5.1 µs', '3.6.0 µs', '3.7.1 µs'],
        ['32×128 / 4 warps', '298.032', '298.224', '298.832'],
        ['32×64 / 8 warps', '436.224', '436.208', '436.976'],
    ], .69, 18)
    a.bullets(s, .88, 4.76, 11.55, [
        '这两个配置中，三个版本的有效 M/N 默认分布均为 vec8，未观察到该发布边界的回归。',
        '本表与第 12 页是不同批次，不能把本表约 436 µs 和另一批约 279 µs 混算收益。',
        '原始 TMA issue 的版本回归与普通路径的长期优化机会，是两种不同证据。',
    ], 18, .21)
    a.note(s, '原始配置的3.6/3.5比值1.00012，95% t区间[0.99882,1.00141]；补充配置0.99989，[0.99973,1.00006]。'
           '3%阈值在测试前冻结，结论仅限一张A10和两个配置。'
           '相同shape的三个release最终TTGIR均核对load和reduction输入编码；order在size1 batch维上的变化不改变有效M/N归属。'
           '这些时间不是今日main的性能测量，也不替代原GB300的TMA验证。')

    s = a.page('附录 C. 参考与证据', '源码版本、原始证据与后续方案索引',
               '源码与数据分别固定版本；所有新增结论均可回到归档文件核对。',
               '随交付：mechanism_audit.md、audit.json、source_audit/、speaker_notes.md')
    refs = [
        ('[1] 原始回归', 'Triton Issue #9703：GB300，TMA descriptor load + partial max', ISSUE),
        ('[2] H100 布局实验', 'results/phase11/stage_b/{summary,interpretation}.md；upstream 20d9d98', None),
        ('[3] A10 两轮实验', 'results/local_a10/ 与 results/local_a10_versions/；本地 commit 942d380', None),
        ('[4] 布局语义', 'Triton / Gluon layouts：逻辑元素到 register / lane / warp 的映射',
         'https://triton-lang.org/main/getting-started/tutorials/gluon/layouts.html'),
        ('[5] 早期代码审计', 'Phase10：32×128/w4 BF16 host 路径 baseline / patch 的原始 PTX', None),
        ('[6] NCU 定义', 'Instructions、Requests、Wavefronts、Bank Conflicts 分别核对',
         'https://docs.nvidia.com/nsight-compute/ProfilingGuide/#shared-memory'),
        ('[7] 后续 review draft', 'phase11/FOLLOWUP_PLAN.md：C1/C2/C3、保留评价与架构迁移', None),
        ('[S1–S5] 当前机制', f'main {MAIN[:12]}：Coalesce、builder、传播与 ThreadLocality', GH + 'lib/Dialect/TritonGPU/Transforms/Coalesce.cpp'),
    ]
    for i, (label, description, url) in enumerate(refs):
        y = 2.05 + i * .59
        a.text(s, .87, y, 2.38, .39, label, 14.5, True)
        a.text(s, 3.27, y, 9.05, .43, description, 14.0, url=url)
    a.note(s, '所有源码文件已随报告归档并记录SHA256。S1为Coalesce.cpp，S2为CoalesceUtils/Utility，S3为TritonGPUAttrDefs.td builder，'
           'S4为RemoveLayoutConversions，S5为OptimizeThreadLocality。详细函数和不可变commit链接见mechanism_audit.md。'
           '报告输入PPTX也保存SHA；旧raw samples、compiler artifacts与NCU reports保持不变。')

    a.save()
    print('Saved editable PPTX:', OUTPUT, '; slides:', len(a.order))


if __name__ == '__main__':
    build()
